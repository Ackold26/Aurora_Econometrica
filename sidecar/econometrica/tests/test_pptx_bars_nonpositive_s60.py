"""s60, п.8(iii) записки MEMO_BACKLOG (INV-50): слайд «полосы mROAS по каналам»
молча отбрасывал каналы с mROAS ≤ 0 – худшая новость портфеля (канал уменьшает
продажи) исчезала с графика, а HTML показывал все каналы.

Стережём на ЗНАЧЕНИЯХ, прочитанных из XML диаграммы готового файла (выгрузка
сохраняется и открывается заново), а не копией логики сборщика:
(а) денежный KPI, два канала ≤ 0 – все 5 категорий, значения = mROAS, порядок
    «лучший сверху» прежний;
(б) штучный KPI (стоимость единицы 1/mROAS при ≤ 0 не определена) – 3 категории
    и сноска с обоими именами;
(в) денежный, все каналы ≤ 0 – диаграмма с отрицательными полосами, без надписи
    M-5 «Ни у одного канала нет положительной отдачи…» (она – только штучный);
(г) режим «доля» – тоже все каналы;
(д) цвет отрицательной полосы – токен semantic.stop, не запасное значение;
(е) invertIfNegative = 0 у ряда и у каждой точки – иначе PowerPoint красит
    отрицательную сплошную полосу в белый;
(ж) HTML и PPTX показывают один и тот же набор каналов в денежном режиме;
плюс: подписи категорий у левого края и запас оси под подпись значения (чтобы
подписи отрицательных полос не лезли на подписи категорий), значение «не
определено» не рисуется нулём, сторож переполнения/контраста чист.
"""
from __future__ import annotations

import copy
import io
import json
import math
import os

import pytest
from pptx import Presentation

from aurora_pptx.builder import AuroraPPTXBuilder
from aurora_pptx.check_overflow import contrast_ratio, guard_summary
from aurora_pptx.tokens import COLORS

_HERE = os.path.dirname(os.path.abspath(__file__))
_FIXTURE = os.path.join(_HERE, "fixtures", "kagocel_builder_payload.json")

_NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
_T_MONEY = "MROAS ПО КАНАЛАМ / МУЛЬТИПЛИКАТОР"
_T_COUNT = "CPU ПО КАНАЛАМ / ₽ ЗА ЕДИНИЦУ"
_T_SHARE = "ДОЛЯ КАНАЛОВ В ЭФФЕКТЕ / %"
_M5 = "Ни у одного канала нет положительной отдачи"
_NOTE_HEAD = "Не показаны:"
_NOTE_NONPOS = "отдача нулевая или отрицательная"
_NOTE_UNDEF = "значение не определено"

_NEG = "OLV"        # mROAS < 0
_ZERO = "Banners"   # mROAS = 0


def _payload(*, mroas: dict | None = None, kind: str | None = None,
             mode: str | None = None) -> dict:
    with open(_FIXTURE, encoding="utf-8") as f:
        p = copy.deepcopy(json.load(f))
    for c in p["channels"]:
        if mroas and c["name"] in mroas:
            c["mroas"] = mroas[c["name"]]
    if kind:
        p["kpi"]["kpi_kind"] = kind
    if mode:
        p["kpi"]["derived_mode"] = mode
    return p


def _two_nonpositive(**kw) -> dict:
    return _payload(mroas={_NEG: -0.35, _ZERO: 0.0}, **kw)


def _share_payload() -> dict:
    # режим «доля»: в поле mroas лежит доля канала (0..1)
    return _payload(
        mroas={"Performance": 0.40, "Social": 0.35, "TRPs бренд (W 25-54)": 0.25,
               _ZERO: 0.0, _NEG: -0.05},
        mode="effectiveness",
    )


def _roundtrip(data: dict):
    """Сборка -> файл в памяти -> открытие заново (как получает клиент)."""
    buf = io.BytesIO()
    AuroraPPTXBuilder(data).build().save(buf)
    buf.seek(0)
    return Presentation(buf), buf.getvalue()


def _slide(prs, title: str):
    for slide in prs.slides:
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        if any(title in t for t in texts):
            return slide, texts
    return None, []


def _bar_chart_xml(slide):
    charts = [sh for sh in slide.shapes if sh.has_chart]
    assert len(charts) == 1, "на слайде полос ровно одна диаграмма"
    return charts[0].chart._chartSpace


def _pts(parent) -> dict[int, str]:
    return {int(pt.get("idx")): pt.findtext("c:v", namespaces=_NS)
            for pt in parent.iterfind(".//c:pt", _NS)}


def _read(cs):
    """(категории, значения) в порядке данных XML (PowerPoint рисует первую ВНИЗУ)."""
    cat = _pts(cs.find(".//c:ser/c:cat", _NS))
    val = _pts(cs.find(".//c:ser/c:val", _NS))
    n = len(cat)
    assert sorted(cat) == list(range(n)) and sorted(val) == list(range(n))
    return [cat[i] for i in range(n)], [float(val[i]) for i in range(n)]


def _display(cs):
    """Порядок сверху вниз: диаграмма хранит данные развёрнутыми."""
    cats, vals = _read(cs)
    return list(reversed(cats)), list(reversed(vals))


def _expected_desc(data: dict) -> list[tuple[str, float]]:
    """Независимо от сборщика: все каналы, лучший сверху (по убыванию)."""
    return sorted(((c["name"], float(c["mroas"])) for c in data["channels"]),
                  key=lambda t: t[1], reverse=True)


def _fill_by_name(cs) -> dict[str, str]:
    cats, _ = _read(cs)
    out = {}
    for dpt in cs.iterfind(".//c:ser/c:dPt", _NS):
        idx = int(dpt.find("c:idx", _NS).get("val"))
        clr = dpt.find("c:spPr/a:solidFill/a:srgbClr", _NS)
        out[cats[idx]] = clr.get("val").upper()
    return out


def _hex(token: str) -> str:
    return token.lstrip("#").upper()


# ─── (а) денежный KPI, два канала ≤ 0 ───────────────────────────────────────

def test_money_two_nonpositive_all_channels_values_and_order():
    data = _two_nonpositive()
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, texts = _slide(prs, _T_MONEY)
    assert slide is not None
    cats, vals = _display(_bar_chart_xml(slide))
    exp = _expected_desc(data)
    assert len(cats) == 5 == len(data["channels"])
    assert cats == [n for n, _ in exp]                       # порядок: лучший сверху
    assert vals == pytest.approx([v for _, v in exp])         # значения = mROAS
    assert {_NEG, _ZERO} <= set(cats)
    assert not any(_NOTE_HEAD in t for t in texts)            # денежный: сноски нет
    assert not any(_M5 in t for t in texts)


def test_money_positive_only_payload_is_unchanged_set_and_order():
    data = _payload()
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, texts = _slide(prs, _T_MONEY)
    cats, vals = _display(_bar_chart_xml(slide))
    exp = _expected_desc(data)
    assert cats == [n for n, _ in exp] and vals == pytest.approx([v for _, v in exp])
    assert not any(_NOTE_HEAD in t for t in texts)


# ─── (б) штучный KPI ────────────────────────────────────────────────────────

def test_count_kpi_drops_nonpositive_and_names_them_in_footnote():
    data = _two_nonpositive(kind="count")
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, texts = _slide(prs, _T_COUNT)
    assert slide is not None
    cats, vals = _display(_bar_chart_xml(slide))
    positive = [(c["name"], float(c["mroas"])) for c in data["channels"]
                if float(c["mroas"]) > 0]
    exp = sorted(((n, 1.0 / v) for n, v in positive), key=lambda t: t[1])  # дешевле – выше
    assert len(cats) == 3
    assert cats == [n for n, _ in exp]
    assert vals == pytest.approx([v for _, v in exp])
    notes = [t for t in texts if _NOTE_HEAD in t]
    assert len(notes) == 1
    assert _NEG in notes[0] and _ZERO in notes[0] and _NOTE_NONPOS in notes[0]
    assert not any(_M5 in t for t in texts)                   # диаграмма есть – M-5 нет


def test_count_kpi_footnote_absent_when_nothing_excluded():
    prs, _ = _roundtrip(_payload(kind="count"))
    slide, texts = _slide(prs, _T_COUNT)
    assert slide is not None and any(sh.has_chart for sh in slide.shapes)
    assert not any(_NOTE_HEAD in t for t in texts)


def test_count_kpi_all_nonpositive_keeps_m5_note_without_chart():
    data = _payload(kind="count")
    for c in data["channels"]:
        c["mroas"] = -0.35
    prs, _ = _roundtrip(data)
    slide, texts = _slide(prs, _T_COUNT)
    assert slide is not None
    assert any(_M5 in t for t in texts)
    assert not any(sh.has_chart for sh in slide.shapes)
    assert not any("безубыточность" in t for t in texts)


# ─── (в) денежный KPI, все каналы ≤ 0 ───────────────────────────────────────

def test_money_all_nonpositive_builds_negative_bars_without_m5_note():
    data = _payload()
    for c in data["channels"]:
        c["mroas"] = -0.35
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, texts = _slide(prs, _T_MONEY)
    assert slide is not None
    cats, vals = _display(_bar_chart_xml(slide))
    assert sorted(cats) == sorted(c["name"] for c in data["channels"])
    assert vals == pytest.approx([-0.35] * 5)
    assert not any(_M5 in t for t in texts)
    assert any("безубыточность" in t for t in texts)          # подпись к диаграмме на месте
    assert not any(_NOTE_HEAD in t for t in texts)


# ─── (г) режим «доля» ───────────────────────────────────────────────────────

def test_share_mode_shows_all_channels_with_percent_labels():
    data = _share_payload()
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, texts = _slide(prs, _T_SHARE)
    assert slide is not None
    cs = _bar_chart_xml(slide)
    cats, vals = _display(cs)
    exp = _expected_desc(data)
    assert cats == [n for n, _ in exp] and len(cats) == 5
    assert vals == pytest.approx([v for _, v in exp])
    fmt = cs.find(".//c:dLbls/c:numFmt", _NS).get("formatCode")
    assert fmt == "0.0%"
    assert not any(_NOTE_HEAD in t for t in texts)
    fills = _fill_by_name(cs)
    assert fills[_NEG] == _hex(COLORS["semantic"]["stop"])


# ─── (д) цвета ──────────────────────────────────────────────────────────────

def test_negative_bar_uses_semantic_stop_token_only_for_negative_values():
    data = _two_nonpositive()
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    fills = _fill_by_name(cs)
    stop = _hex(COLORS["semantic"]["stop"])
    gold = _hex(COLORS["brand"]["gold"]["primary"])
    muted = _hex(COLORS["brand"]["deep"]["40"])
    best = _expected_desc(data)[0][0]
    assert fills[_NEG] == stop
    assert fills[_ZERO] != stop and fills[_ZERO] == muted      # ноль – не «стоп»
    assert fills[best] == gold                                 # лучший (> 0) – золото
    for name, fill in fills.items():
        if name not in (_NEG, best):
            assert fill == muted and fill != stop
    # отрицательная полоса читается на белом (графика: не ниже 3:1)
    assert contrast_ratio(stop, "FFFFFF") >= 3.0


def test_best_channel_is_not_gold_when_all_nonpositive():
    data = _payload()
    for c in data["channels"]:
        c["mroas"] = -0.35
    prs, _ = _roundtrip(data)
    slide, _t = _slide(prs, _T_MONEY)
    fills = _fill_by_name(_bar_chart_xml(slide))
    gold = _hex(COLORS["brand"]["gold"]["primary"])
    assert gold not in fills.values()
    assert set(fills.values()) == {_hex(COLORS["semantic"]["stop"])}


def test_zero_best_channel_is_not_gold():
    # лучший канал = 0.0, остальные отрицательные: «героя» золотом не красим
    data = _payload()
    for c in data["channels"]:
        c["mroas"] = 0.0 if c["name"] == _ZERO else -0.35
    prs, _ = _roundtrip(data)
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    fills = _fill_by_name(cs)
    gold = _hex(COLORS["brand"]["gold"]["primary"])
    assert gold not in fills.values()
    assert fills[_ZERO] == _hex(COLORS["brand"]["deep"]["40"])
    assert [v for n, v in fills.items() if n != _ZERO] == [_hex(COLORS["semantic"]["stop"])] * 4


def test_data_label_color_readable_on_white():
    prs, _ = _roundtrip(_two_nonpositive())
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    clr = cs.find(".//c:dLbls/c:txPr//a:solidFill/a:srgbClr", _NS).get("val")
    assert contrast_ratio(clr, "FFFFFF") >= 4.5


# ─── (е) invertIfNegative ───────────────────────────────────────────────────

def test_invert_if_negative_is_false_on_series_and_every_point():
    prs, _ = _roundtrip(_two_nonpositive())
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    ser_flags = [e.get("val") for e in cs.iterfind(".//c:ser/c:invertIfNegative", _NS)]
    assert ser_flags == ["0"]
    dpts = cs.findall(".//c:ser/c:dPt", _NS)
    assert len(dpts) == 5
    for dpt in dpts:
        assert [e.get("val") for e in dpt.iterfind("c:invertIfNegative", _NS)] == ["0"]
        # порядок схемы CT_DPt: idx, invertIfNegative, …, spPr
        tags = [el.tag.split("}")[1] for el in dpt]
        assert tags.index("idx") < tags.index("invertIfNegative") < tags.index("spPr")


# ─── (ж) HTML и PPTX – один набор каналов ───────────────────────────────────

def test_html_and_pptx_show_the_same_channels_in_money_mode():
    from aurora_html.builder import AuroraHTMLBuilder
    data = _two_nonpositive()
    html_names = json.loads(
        AuroraHTMLBuilder(copy.deepcopy(data))._chart_data_json())["mroas"]["names"]
    prs, _ = _roundtrip(copy.deepcopy(data))
    slide, _t = _slide(prs, _T_MONEY)
    cats, _v = _display(_bar_chart_xml(slide))
    assert len(html_names) == 5
    assert cats == html_names


# ─── подписи значений не налезают на подписи категорий ──────────────────────

def test_axis_leaves_room_for_negative_labels_and_category_labels_stay_left():
    prs, _ = _roundtrip(_two_nonpositive())
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    lo = float(cs.find(".//c:valAx/c:scaling/c:min", _NS).get("val"))
    hi = float(cs.find(".//c:valAx/c:scaling/c:max", _NS).get("val"))
    _c, vals = _read(cs)
    assert lo < min(vals) < 0
    assert (min(vals) - lo) >= 0.10 * (hi - lo)               # место под подпись значения
    assert cs.find(".//c:catAx/c:tickLblPos", _NS).get("val") == "low"


def test_positive_only_chart_keeps_zero_axis_and_default_label_position():
    prs, _ = _roundtrip(_payload())
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    assert float(cs.find(".//c:valAx/c:scaling/c:min", _NS).get("val")) == 0.0
    assert cs.find(".//c:catAx/c:tickLblPos", _NS).get("val") != "low"


def test_all_nonpositive_axis_is_compact_and_contains_data():
    data = _payload()
    for c in data["channels"]:
        c["mroas"] = -0.35
    prs, _ = _roundtrip(data)
    slide, _t = _slide(prs, _T_MONEY)
    cs = _bar_chart_xml(slide)
    lo = float(cs.find(".//c:valAx/c:scaling/c:min", _NS).get("val"))
    hi = float(cs.find(".//c:valAx/c:scaling/c:max", _NS).get("val"))
    assert lo < -0.35 and 0 < hi <= 1.0


def test_guard_overflow_and_contrast_clean_on_export(tmp_path):
    cases = {
        "money": _two_nonpositive(),
        "count": _two_nonpositive(kind="count"),
        "share": _share_payload(),
    }
    allneg = _payload()
    for c in allneg["channels"]:
        c["mroas"] = -0.35
    cases["money_all"] = allneg
    for name, data in cases.items():
        _prs, raw = _roundtrip(data)
        path = tmp_path / f"{name}.pptx"
        path.write_bytes(raw)
        bad, line, details = guard_summary(str(path))
        assert not bad, (name, line, details)


# ─── значение «не определено» не рисуется нулём ─────────────────────────────

@pytest.mark.parametrize("bad_value", [None, float("nan")], ids=["none", "nan"])
def test_undefined_value_is_not_drawn_as_zero_and_is_named(bad_value):
    data = _payload(mroas={_NEG: bad_value})
    if bad_value is None:
        for c in data["channels"]:
            if c["name"] == _NEG:
                c["mroas"] = None
    prs, _ = _roundtrip(data)
    slide, texts = _slide(prs, _T_MONEY)
    cats, vals = _display(_bar_chart_xml(slide))
    assert _NEG not in cats and len(cats) == 4
    assert all(math.isfinite(v) for v in vals)
    notes = [t for t in texts if _NOTE_HEAD in t]
    assert len(notes) == 1 and _NEG in notes[0] and _NOTE_UNDEF in notes[0]
    assert _NOTE_NONPOS not in notes[0]
