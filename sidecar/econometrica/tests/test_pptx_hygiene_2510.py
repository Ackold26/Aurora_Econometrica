"""Гигиена PPTX (s59, блок P3, 2.5.10): контраст золотого текста, сторож после сборки,
правила сторожа, герой «опережает» не «тянет вниз», оси столбчатых диаграмм.

H1  Золотой primary как ТЕКСТ на светлом (2,36:1) заменён на gold_muted (#8C7142, 4,61:1 на белом),
    три надписи на плашке bg_quiet – на производный gold_on_quiet (#826C48, 4,60:1 на #F7F5EE).
    Линии, заливки, полоса героя остаются золотым primary. Cut и Reduce в таблице действий
    различаются цветом, оба >= 4,5:1.
H2  `engines.pptx_export.build_pptx` после штампа вызывает сторож в СВОЁМ try: одна запись WARNING,
    сбой сторожа не рушит выгрузку, файл и его слайды те же.
H3  `check_overflow.check_hygiene` / `guard_summary`: контраст (гейт «0» на выгрузке) и реестр кеглей.
H4  Герой, названный «опережает», не входит в «тянут портфель вниз» – одинаково в HTML и PPTX;
    ОТВЕТ, таблица и остальной текст отчёта не меняются.
H5  У BAR_*/COLUMN_* с неотрицательными данными ось значений начинается с 0 (или авто), максимум >= данных.

Мутации – Projects/_s59/FIX3_hygiene_2510.md.
"""
from __future__ import annotations

import copy
import io
import json
import os
import re
import sys
import zipfile
from functools import lru_cache

import pytest

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))
sys.path.insert(0, _HERE)

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import test_underperformers_unified as U  # noqa: E402
from aurora_pptx.builder import AuroraPPTXBuilder  # noqa: E402

try:
    from econometrica.aurora_pptx import check_overflow as guard
except ImportError:
    from aurora_pptx import check_overflow as guard

_FIXTURE = os.path.join(_HERE, "fixtures", "kagocel_builder_payload.json")
_WHITE = "FFFFFF"


# ─── независимый расчёт контраста (WCAG 2.x), не из кода сторожа ─────────────

def _lum(hex6: str) -> float:
    ch = []
    for i in (0, 2, 4):
        c = int(hex6[i:i + 2], 16) / 255.0
        ch.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]


def _cr(a: str, b: str) -> float:
    la, lb = _lum(a), _lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def _tokens():
    tokens = pytest.importorskip("aurora_tokens")
    brand = tokens.COLORS["brand"]
    return {
        "gold": brand["gold"]["primary"].lstrip("#").upper(),
        "gold_muted": brand["gold"]["muted"].lstrip("#").upper(),
        "quiet": brand["bg"]["quiet"].lstrip("#").upper(),
        "white": brand["bg"]["white"].lstrip("#").upper(),
        "deep_100": brand["deep"]["100"].lstrip("#").upper(),
    }


# ─── колоды: фикстура, полная нагрузка, варианты надёжности ─────────────────

_WF = {"labels": ["База", "TRPs бренд (W 25-54)", "Social", "OLV", "Banners", "Performance", "Итого"],
       "values": [600.0, 50.0, 40.0, 20.0, 10.0, 5.0, 725.0],
       "types": ["base", "pos", "pos", "pos", "pos", "pos", "total"]}


def _fixture_payload() -> dict:
    with open(_FIXTURE, encoding="utf-8") as fh:
        return json.load(fh)


def _full_payload() -> dict:
    """Фикстура + все вставные слайды (по образцу scan.py аудита PLANAUDIT_P3_2510)."""
    d = _fixture_payload()
    d["waterfall"] = _WF
    d["backtest"] = {"status": "ok", "windows_hit_total": 3, "windows_with_interval": 4,
                     "granularity": "M", "horizon_periods": 3,
                     "windows": [{"window": "2026-01", "actual_total": 100, "predicted_total": 95,
                                  "pi_low_total": 80, "pi_high_total": 110, "hit_total": True}],
                     "mape_naive_best": 12.0, "mape_model": 8.0,
                     "verdict_text": "Модель точнее наивного прогноза."}
    d["generation_compare"] = {"status": "ok", "summary": {"headline": "ROI каналов пересчитан."},
                               "channels": [{"name": "Social", "roi_old": 1.2, "roi_ci_old": [1.0, 1.4],
                                             "roi_new": 1.5, "roi_ci_new": [1.3, 1.7], "verdict_ru": "вырос"}]}
    d["promises_summary"] = {"kept": 2, "missed": 1, "examples": [
        {"action_text": "Увеличить Social", "status": "kept", "status_ru": "сбылось"}]}
    d["forecast"] = {"status": "ok", "scenarios": [{
        "name": "Базовый", "variant_id": "v1", "total_kpi": 460.0, "total_kpi_ci_low": 414.0,
        "total_kpi_ci_high": 506.0, "total_spend_money": 5_000_000.0, "roas_money": 0.092}],
        "accepted_variant": "v1", "disclaimers": ["Прогноз при неизменных условиях рынка"]}
    n = 20
    d.setdefault("diagnostics", {})["actual_vs_predicted"] = {
        "dates": [f"2025-{(i % 12) + 1:02d}-01" for i in range(n)],
        "actual": [100 + i for i in range(n)], "predicted": [98 + i for i in range(n)],
        "ci_low": [90 + i for i in range(n)], "ci_high": [110 + i for i in range(n)]}
    return d


def _deck_payloads() -> dict:
    out = {"fixture": _fixture_payload(), "full": _full_payload()}
    for verdict in ("reliable", "unreliable", "refused", "unknown"):
        d = _full_payload()
        d["diagnostics"]["honesty_verdict"] = verdict
        out[f"full_{verdict}"] = d
    missed = _full_payload()
    missed["promises_summary"]["examples"] = [
        {"action_text": "Сократить TRPs", "status": "missed", "status_ru": "не сбылось"}]
    out["full_promise_missed"] = missed
    return out


@pytest.fixture(scope="module")
def decks(tmp_path_factory):
    """{имя: путь к сохранённому .pptx} – колоды штатной сборкой."""
    base = tmp_path_factory.mktemp("hygiene_decks")
    paths = {}
    for name, payload in _deck_payloads().items():
        prs = AuroraPPTXBuilder(copy.deepcopy(payload)).build()
        p = str(base / f"{name}.pptx")
        prs.save(p)
        paths[name] = p
    return paths


_DECK_NAMES = ["fixture", "full", "full_reliable", "full_unreliable", "full_refused",
               "full_unknown", "full_promise_missed"]


def _runs(prs):
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            yield si, sh, r


def _run_hex(r):
    try:
        if r.font.color is not None and r.font.color.type is not None:
            return str(r.font.color.rgb)
    except Exception:
        pass
    return None


# ═══ H1 ═════════════════════════════════════════════════════════════════════

def test_h1_token_contrast_and_builder_colors():
    t = _tokens()
    b = AuroraPPTXBuilder(_fixture_payload())
    assert _cr(t["gold"], t["white"]) < 3.0                      # зачем правка: primary не читается
    assert str(b.gold_muted) == t["gold_muted"]
    assert _cr(t["gold_muted"], t["white"]) >= 4.5
    assert str(b.bg_quiet) == t["quiet"]
    assert _cr(str(b.gold_on_quiet), t["quiet"]) >= 4.5
    assert _cr(str(b.gold_on_quiet), t["white"]) >= 4.5
    assert str(b.gold) == t["gold"]                               # бренд-токен не тронут


def test_h1_fallback_colors_without_generated_tokens_stay_readable(monkeypatch):
    """aurora_tokens.py генерируется сборкой (в .gitignore): без него builder берёт запасные значения.
    Они обязаны давать те же читаемые цвета (раньше gold_muted был #B8975D, 2,75:1, а плашка #F7F7F7)."""
    import aurora_pptx.builder as bmod
    monkeypatch.setattr(bmod, "COLORS", {"navy_deep": "#0A1628", "gold": "#C5A46D"})   # плоский запасной набор
    b = AuroraPPTXBuilder(_fixture_payload())
    assert _cr(str(b.gold_muted), _WHITE) >= 4.5, str(b.gold_muted)
    assert str(b.bg_quiet) == "F7F5EE", str(b.bg_quiet)
    assert _cr(str(b.gold_on_quiet), str(b.bg_quiet)) >= 4.5


@pytest.mark.parametrize("name", _DECK_NAMES)
def test_h1_no_text_run_uses_gold_primary(name, decks):
    """Золотой primary в шрифте не встречается нигде в колоде (все ветки: вставные слайды,
    варианты надёжности, «не сбылось»). Независимый обход XML, не через сторож."""
    gold = _tokens()["gold"]
    bad = [(si, r.text[:30]) for si, _, r in _runs(Presentation(decks[name])) if _run_hex(r) == gold]
    assert bad == [], bad


@pytest.mark.parametrize("name", _DECK_NAMES)
def test_h1_every_run_passes_contrast_independent(name, decks):
    """Каждый прогон с явным цветом: >= 4,5 (крупный >= 3) к своему фону – белый или плашка.
    Расчёт и поиск фона – свои, не из сторожа."""
    t = _tokens()
    prs = Presentation(decks[name])
    bad = []
    for si, slide in enumerate(prs.slides, 1):
        plates = []
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    if sh.fill.type == 1 and sh.width > Inches(0.3) and sh.height > Inches(0.3):
                        plates.append((sh.left, sh.top, sh.left + sh.width, sh.top + sh.height,
                                       str(sh.fill.fore_color.rgb)))
                except Exception:
                    pass
        for sh in slide.shapes:
            if not (getattr(sh, "has_text_frame", False) and sh.has_text_frame):
                continue
            cx, cy = sh.left + sh.width // 2, sh.top + sh.height // 2
            cands = [(x1 - x0) * (y1 - y0) for x0, y0, x1, y1, _ in plates if x0 <= cx <= x1 and y0 <= cy <= y1]
            bg = t["white"]
            if cands:
                bg = min((p for p in plates if p[0] <= cx <= p[2] and p[1] <= cy <= p[3]),
                         key=lambda p: (p[2] - p[0]) * (p[3] - p[1]))[4]
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    fg = _run_hex(r)
                    txt = r.text.strip()
                    if not txt or fg is None or all(c in ".· \t…" for c in txt):
                        continue
                    size = r.font.size.pt if r.font.size else 12.0
                    large = size >= 18 or (size >= 14 and bool(r.font.bold))
                    if _cr(fg, bg) < (3.0 if large else 4.5):
                        bad.append((si, size, fg, bg))
    assert bad == [], bad


def test_h1_lines_fills_and_hero_bar_stay_gold_primary(decks):
    """Подмена идёт только в тексте: полоса-заливка, линии и точка героя на диаграмме – золотой primary."""
    gold = _tokens()["gold"]
    prs = Presentation(decks["full"])
    fills = lines = hero_points = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == gold:
                        fills += 1
                except Exception:
                    pass
            if sh.shape_type == MSO_SHAPE_TYPE.LINE:
                try:
                    if str(sh.line.color.rgb) == gold:
                        lines += 1
                except Exception:
                    pass
            if getattr(sh, "has_chart", False) and sh.has_chart and sh.chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
                ser = sh.chart.plots[0].series[0]
                for i in range(len(list(ser.values))):
                    try:
                        if str(ser.points[i].format.fill.fore_color.rgb) == gold:
                            hero_points += 1
                    except Exception:
                        pass
    assert fills >= 1 and lines >= 5 and hero_points == 1, (fills, lines, hero_points)


def _verdict_runs(payload):
    prs = AuroraPPTXBuilder(copy.deepcopy(payload)).build()
    got = {}
    for si, _, r in _runs(prs):
        for key in ("Остановить", "Сократить"):
            if r.text.startswith(key) and r.font.bold and r.font.size and r.font.size.pt == 11:
                got[key] = _run_hex(r)
    return got


def test_h1_ink_replaces_gold_primary_in_all_three_text_helpers():
    """Золото в тексте подменяется в каждом помощнике, который пишет цвет шрифта: _text, _rich, _paragraphs.
    Прямая проверка – не каждый путь исполняется в колодах-фикстурах."""
    t = _tokens()
    b = AuroraPPTXBuilder(_fixture_payload())
    slide = b._blank()
    b._text(slide, 1, 1, 3, 0.3, "t", color=b.gold)
    b._rich(slide, 1, 2, 3, 0.3, [("r", {"color": b.gold})])
    b._paragraphs(slide, 1, 3, 3, 0.3, [("p", {"color": b.gold})])
    b._text(slide, 1, 4, 3, 0.3, "q", color=b.gold_on_quiet)
    got = {r.text: _run_hex(r) for _, _, r in _runs_of_slide(slide)}
    assert got == {"t": t["gold_muted"], "r": t["gold_muted"], "p": t["gold_muted"],
                   "q": str(b.gold_on_quiet)}, got


def _runs_of_slide(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    yield 1, sh, r


def test_h3b_cut_and_reduce_differ_in_color_and_both_readable():
    payload = U._scenario_data(spec={"Banners": {"action": "Reduce"}})
    got = _verdict_runs(payload)
    assert set(got) == {"Остановить", "Сократить"}, got
    assert got["Остановить"] != got["Сократить"], got
    for word, hex6 in got.items():
        assert _cr(hex6, _WHITE) >= 4.5, (word, hex6, _cr(hex6, _WHITE))


# ═══ H3: правила сторожа ════════════════════════════════════════════════════

def _synthetic_deck(path, runs, plate=None):
    """runs = [(текст, кегль, жирный, цвет-hex)]; plate – hex заливки плашки под всеми текстами."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if plate:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(8), Inches(5))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(plate)
    y = 0.8
    for text, size, bold, color in runs:
        tb = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(6), Inches(0.4))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        y += 0.6
    prs.save(path)
    return path


def test_h3_wcag_formula_reference_values():
    assert guard.contrast_ratio("000000", "FFFFFF") == pytest.approx(21.0, rel=1e-6)
    assert guard.contrast_ratio("C5A46D", "FFFFFF") == pytest.approx(2.36, abs=0.01)


def test_h3_contrast_rule_finds_planted_violations_and_stays_quiet_on_clean(tmp_path):
    gold, muted, ink = "C5A46D", "8C7142", "0A1628"
    p = _synthetic_deck(str(tmp_path / "bad.pptx"), [
        ("золото обычное", 10, False, gold),       # 2,36 < 4,5
        ("золото крупное", 20, False, gold),       # 2,36 < 3
        ("тёмное чистое", 10, False, ink),
        ("золото приглушённое на белом", 10, False, muted),   # 4,61 – проходит
        ("золото крупное жирное 14", 14, True, "999999"),       # 2,85 < 3
        ("......", 10, False, "DDDDDD"),            # отточие – не считается
    ])
    hyg = guard.check_hygiene(p)
    assert len(hyg["contrast"]) == 3, hyg["contrast"]
    clean = _synthetic_deck(str(tmp_path / "clean.pptx"), [
        ("чисто 1", 10, False, ink), ("чисто 2", 12, False, muted), ("крупное", 20, False, "8C7142")])
    assert guard.check_hygiene(clean)["contrast"] == []


def test_h3_contrast_rule_uses_plate_background(tmp_path):
    """gold_muted на белом проходит (4,61), на плашке #F7F5EE – нет (4,23): фон берётся из плашки."""
    runs = [("надпись", 10, False, "8C7142")]
    on_white = _synthetic_deck(str(tmp_path / "w.pptx"), runs)
    on_plate = _synthetic_deck(str(tmp_path / "p.pptx"), runs, plate="F7F5EE")
    assert guard.check_hygiene(on_white)["contrast"] == []
    assert len(guard.check_hygiene(on_plate)["contrast"]) == 1


def test_h3_large_text_uses_the_3_to_1_threshold(tmp_path):
    """Крупный текст (>= 18 pt или >= 14 pt жирный) судится по 3:1: gold_muted на плашке (4,23) для него проходит,
    для обычного размера – нет."""
    plate = "F7F5EE"
    big = _synthetic_deck(str(tmp_path / "big.pptx"), [("a", 20, False, "8C7142"), ("b", 14, True, "8C7142")],
                          plate=plate)
    assert guard.check_hygiene(big)["contrast"] == []
    small = _synthetic_deck(str(tmp_path / "small.pptx"), [("a", 13, True, "8C7142"), ("b", 17, False, "8C7142")],
                            plate=plate)
    assert len(guard.check_hygiene(small)["contrast"]) == 2


def test_h3_size_registry_counts_runs_below_9_and_below_10(tmp_path):
    p = _synthetic_deck(str(tmp_path / "s.pptx"), [
        ("a", 8, False, "0A1628"), ("b", 9.5, False, "0A1628"), ("c", 12, False, "0A1628")])
    hyg = guard.check_hygiene(p)
    assert hyg["small9"] == {1: 1} and hyg["small10"] == {1: 2}, hyg


def test_h3_guard_summary_one_line_without_client_text(tmp_path):
    p = _synthetic_deck(str(tmp_path / "g.pptx"), [("СЕКРЕТНЫЙ ТЕКСТ КЛИЕНТА", 8, False, "C5A46D")])
    has, line, details = guard.guard_summary(p)
    assert has is True and "\n" not in line
    assert "контраст ниже порога=1" in line and "<9 pt: 1" in line, line
    assert "СЕКРЕТНЫЙ" not in line and not any("СЕКРЕТНЫЙ" in d for d in details), (line, details)


def test_h3_guard_details_name_slide_and_kind_but_never_client_text(tmp_path):
    """Переполнение: у самого сторожа в подробностях идёт кусок текста бокса; в журнал – только слайд и вид."""
    runs = [(f"заполнитель {i}", 8, False, "0A1628") for i in range(10)]
    runs.append(("СЕКРЕТНЫЙ ТЕКСТ КЛИЕНТА ВНИЗУ СЛАЙДА", 40, False, "0A1628"))
    p = _synthetic_deck(str(tmp_path / "over.pptx"), runs)
    issues, _ = guard.check(p)
    assert any(k == "OVERFLOW" for _, k, _ in issues), issues
    assert any("СЕКРЕТНЫЙ" in d for _, _, d in issues)          # сторож сам текст знает...
    has, line, details = guard.guard_summary(p)
    assert has is True and "переполнение=0" not in line, line
    assert any("OVERFLOW" in d for d in details), details
    assert not any("СЕКРЕТНЫЙ" in d for d in details) and "СЕКРЕТНЫЙ" not in line   # ...а журнал – нет


@pytest.mark.parametrize("name", _DECK_NAMES)
def test_h3_contrast_gate_zero_on_real_decks(name, decks):
    """Гейт: на выгрузке штатной сборки правило контраста сторожа даёт 0."""
    assert guard.check_hygiene(decks[name])["contrast"] == []


@pytest.mark.parametrize("name", ["fixture", "full"])
def test_h3_overflow_still_clean_on_real_decks(name, decks):
    issues, _total = guard.check(decks[name])
    assert issues == [], issues


# ═══ H2: сторож в боевой выгрузке ═══════════════════════════════════════════

def _export(tmp_path, monkeypatch, name="out.pptx"):
    import engines.pptx_export as px
    payload = _fixture_payload()
    monkeypatch.setattr(px, "_map_pipeline_to_builder_data", lambda *a, **k: copy.deepcopy(payload))
    return px.build_pptx({}, {}, {}, str(tmp_path / name))


def _slide_xml(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in sorted(z.namelist()) if n.startswith("ppt/slides/slide")}


def test_h2_export_logs_one_info_line_and_returns_ok(tmp_path, monkeypatch, caplog):
    """Штатная выгрузка: из нарушений только реестр мелкого шрифта – одна запись уровня INFO (не WARNING)."""
    import logging
    caplog.set_level(logging.DEBUG, logger="econometrica")
    res = _export(tmp_path, monkeypatch)
    assert res["status"] == "ok" and os.path.exists(res["path"]), res
    lines = [r for r in caplog.records if r.getMessage().startswith("pptx_guard:") and r.levelno >= logging.INFO]
    assert len(lines) == 1 and lines[0].levelno == logging.INFO, [r.getMessage() for r in lines]
    msg = lines[0].getMessage()
    assert "слайдов=12" in msg and "контраст ниже порога=0" in msg, msg
    assert "переполнение=0, наложение=0" in msg and "<9 pt: 0" not in msg, msg   # реестр шрифта непуст, уровень не поднял
    for r in caplog.records:                                     # ни в одной записи нет текста клиента
        assert "Kagocel" not in r.getMessage() or r.getMessage().startswith("build_pptx OK")
    assert "Social" not in msg and "TRPs" not in msg


def test_h2_guard_failure_does_not_break_export_and_slides_are_the_same(tmp_path, monkeypatch, caplog):
    import logging
    ok = _export(tmp_path, monkeypatch, "with_guard.pptx")

    def boom(_path):
        raise RuntimeError("сторож сломан")

    monkeypatch.setattr(guard, "guard_summary", boom)
    caplog.clear()
    caplog.set_level(logging.DEBUG, logger="econometrica")
    res = _export(tmp_path, monkeypatch, "broken_guard.pptx")
    assert res["status"] == "ok", res
    assert os.path.exists(res["path"]) and res["slides"] == ok["slides"]
    warns = [r for r in caplog.records if r.levelno == logging.WARNING and "сторож" in r.getMessage()]
    assert len(warns) == 1, [r.getMessage() for r in caplog.records]
    assert not [r for r in caplog.records if r.getMessage().startswith("pptx_guard:")]
    assert _slide_xml(ok["path"]) == _slide_xml(res["path"])     # слайды те же со сторожем и без


def test_h2_details_go_to_debug_not_warning(tmp_path, monkeypatch, caplog):
    """Одна строка WARNING на выгрузку: подробности по нарушениям – уровнем DEBUG."""
    import logging
    monkeypatch.setattr(guard, "guard_summary",
                        lambda _p: (True, "pptx_guard: итог", ["слайд 1 CONTRAST 2.36:1", "слайд 2 OVERFLOW"]))
    caplog.set_level(logging.DEBUG, logger="econometrica")
    res = _export(tmp_path, monkeypatch)
    assert res["status"] == "ok"
    guard_recs = [r for r in caplog.records if r.getMessage().startswith("pptx_guard:")]
    assert [(r.levelno, r.getMessage()) for r in guard_recs] == [
        (logging.WARNING, "pptx_guard: итог"),
        (logging.DEBUG, "pptx_guard: слайд 1 CONTRAST 2.36:1"),
        (logging.DEBUG, "pptx_guard: слайд 2 OVERFLOW")]


# ─── уровень записи итога: INFO – только реестр мелкого шрифта, WARNING – настоящие нарушения ───

def _export_synthetic(tmp_path, monkeypatch, runs, name="syn.pptx"):
    """Боевая выгрузка (`build_pptx` -> штамп -> сторож), но сборщик отдаёт заранее заданную колоду."""
    import aurora_pptx
    import engines.pptx_export as px
    src = _synthetic_deck(str(tmp_path / ("src_" + name)), runs)
    monkeypatch.setattr(aurora_pptx, "build_pptx", lambda data, lang="ru": Presentation(src))
    monkeypatch.setattr(px, "_map_pipeline_to_builder_data", lambda *a, **k: {})
    return px.build_pptx({}, {}, {}, str(tmp_path / name))


def _guard_records(caplog):
    import logging
    return [r for r in caplog.records
            if r.getMessage().startswith("pptx_guard:") and r.levelno >= logging.INFO]


def test_h2_level_info_when_only_small_font_registry(tmp_path, monkeypatch, caplog):
    import logging
    caplog.set_level(logging.DEBUG, logger="econometrica")
    res = _export_synthetic(tmp_path, monkeypatch, [("мелкий 8", 8, False, "0A1628"), ("мелкий 9,5", 9.5, False, "0A1628")])
    assert res["status"] == "ok", res
    recs = _guard_records(caplog)
    assert len(recs) == 1 and recs[0].levelno == logging.INFO, [(r.levelno, r.getMessage()) for r in recs]
    msg = recs[0].getMessage()
    assert "переполнение=0, наложение=0" in msg and "контраст ниже порога=0" in msg, msg
    assert "<9 pt: 1" in msg and "<10 pt: 2" in msg, msg          # реестр в строке остался, формат прежний


def test_h2_level_warning_when_overflow(tmp_path, monkeypatch, caplog):
    import logging
    caplog.set_level(logging.DEBUG, logger="econometrica")
    runs = [(f"заполнитель {i}", 8, False, "0A1628") for i in range(10)]
    runs.append(("низ слайда", 40, False, "0A1628"))
    res = _export_synthetic(tmp_path, monkeypatch, runs)
    assert res["status"] == "ok", res
    recs = _guard_records(caplog)
    assert len(recs) == 1 and recs[0].levelno == logging.WARNING, [(r.levelno, r.getMessage()) for r in recs]
    msg = recs[0].getMessage()
    assert "переполнение=0" not in msg and "контраст ниже порога=0" in msg, msg


def test_h2_level_warning_when_contrast_below_threshold(tmp_path, monkeypatch, caplog):
    import logging
    caplog.set_level(logging.DEBUG, logger="econometrica")
    res = _export_synthetic(tmp_path, monkeypatch, [("золото обычное", 10, False, "C5A46D")])
    assert res["status"] == "ok", res
    recs = _guard_records(caplog)
    assert len(recs) == 1 and recs[0].levelno == logging.WARNING, [(r.levelno, r.getMessage()) for r in recs]
    msg = recs[0].getMessage()
    assert "контраст ниже порога=1" in msg and "переполнение=0, наложение=0" in msg, msg


def test_h3_guard_summary_flag_ignores_small_font_registry(tmp_path):
    """Первое значение guard_summary – «настоящие» нарушения: реестр мелкого шрифта его не поднимает."""
    only_small = _synthetic_deck(str(tmp_path / "small.pptx"), [("a", 8, False, "0A1628")])
    has, line, _ = guard.guard_summary(only_small)
    assert has is False and "<9 pt: 1" in line, (has, line)
    clean = _synthetic_deck(str(tmp_path / "clean.pptx"), [("a", 12, False, "0A1628")])
    assert guard.guard_summary(clean)[0] is False


# ═══ H4: «опережает» и «тянут» ══════════════════════════════════════════════

def _build_both(spec=None, **kw):
    from aurora_html import build_html
    data = U._scenario_data(spec=spec, **kw)
    page = build_html(copy.deepcopy(data))
    prs = AuroraPPTXBuilder(copy.deepcopy(data)).build()
    blocks, _ = U._pptx_blocks(prs)
    return data, U._html_blocks(page), blocks, page, U._all_text(prs)


_HERO = "Social"
_HERO_PART = "По mROAS Social опережает (27.03×)."
# s60 (г): у героя с действием Cut/Reduce часть «опережает» идёт с хвостом (engines.channel_action.HERO_CUT_CLAUSE)
_HERO_PART_CUT = "По mROAS Social опережает (27.03×), но оптимизатор рекомендует сократить бюджет канала."


def test_h4_hero_not_leader_with_reduce_is_dropped_from_pull_list_in_both_formats():
    """Герой – Reduce с mROAS 27 (насыщение), лидер – Performance: «опережает» печатается (с хвостом s60:
    оптимизатор рекомендует сократить канал), из «тянут» герой убран, глагол пересчитан: «TRPs ... тянет»
    (один), а не «тянут»."""
    data, html, pptx, *_ = _build_both(spec={_HERO: {"action": "Reduce"}})
    assert data["narrative_facts"]["underperformer_names_all"] == [U._TRPS, _HERO]   # факт адаптера не меняется
    for text in (html["ЗАТРУДНЕНИЕ"], pptx["ПРОБЛЕМА"]):
        assert _HERO_PART_CUT in text, text
        assert U._pull(text) == ([U._TRPS], "тянет"), text


def test_h4_hero_as_cut_in_all_cut_is_dropped_in_both_formats():
    _, html, pptx, *_ = _build_both(balanced=True, spec={
        n: {"action": "Cut"} for n in ("Performance", "Social", "Banners", "OLV")})
    for text in (html["ЗАТРУДНЕНИЕ"], pptx["ПРОБЛЕМА"]):
        names, verb = U._pull(text)
        assert _HERO not in names and verb == "тянут" and len(names) == 4, text
        assert _HERO_PART_CUT in text


def test_h4_hero_that_is_leader_is_not_dropped():
    """Герой = лидер: часть «опережает» не печатается – исключать нечего, герой остаётся в «тянут»."""
    data, html, pptx, *_ = _build_both(balanced=True, spec={"Performance": {"action": "Reduce", "mroas": 40.0}})
    f = data["narrative_facts"]
    assert f["hero_channel"] == f["leader_channel"] == "Performance"
    for text in (html["ЗАТРУДНЕНИЕ"], pptx["ПРОБЛЕМА"]):
        assert "опережает" not in text, text
        assert "Performance" in U._pull(text)[0], text


def test_h4_hero_only_underperformer_pull_part_disappears_and_html_not_fallback():
    """Единственный отстающий – герой: части «тянут» нет; HTML не уходит в запасную фразу (есть «опережает»)."""
    _, html, pptx, *_ = _build_both(balanced=True, spec={U._TRPS: {"action": "Hold"}, _HERO: {"action": "Reduce"}})
    assert html["ЗАТРУДНЕНИЕ"] == _HERO_PART_CUT, html["ЗАТРУДНЕНИЕ"]
    assert html["ЗАТРУДНЕНИЕ"] != U._FALLBACK
    assert pptx["ПРОБЛЕМА"].startswith(_HERO_PART_CUT), pptx["ПРОБЛЕМА"]
    for text in (html["ЗАТРУДНЕНИЕ"], pptx["ПРОБЛЕМА"]):
        assert "тянет" not in text and "тянут" not in text, text


def test_h4_only_the_problem_text_changes_answer_and_table_are_verbatim(monkeypatch):
    """С правкой и без неё (помощник = тождественный список) отчёты отличаются ТОЛЬКО абзацем
    ПРОБЛЕМА/ЗАТРУДНЕНИЕ: ОТВЕТ, таблица действий и остальной текст дословно прежние."""
    import engines.channel_action as ca
    spec = {_HERO: {"action": "Reduce"}}
    _, html_new, pptx_new, page_new, all_new = _build_both(spec=spec)
    monkeypatch.setattr(ca, "drop_named_hero", lambda names, hero, hero_named: list(names))
    _, html_old, pptx_old, page_old, all_old = _build_both(spec=spec)
    assert html_new["ЗАТРУДНЕНИЕ"] != html_old["ЗАТРУДНЕНИЕ"]
    assert pptx_new["ПРОБЛЕМА"] != pptx_old["ПРОБЛЕМА"]
    for label in ("ОТВЕТ", "СИТУАЦИЯ", "ВОПРОС"):
        assert html_new.get(label) == html_old.get(label), label
    assert pptx_new["ОТВЕТ"] == pptx_old["ОТВЕТ"]
    # время сборки в странице (чч:мм, чч:мм:сс) между двумя сборками может смениться – не сравниваем
    _time = re.compile(r"\b\d{2}:\d{2}(?::\d{2})?\b")
    assert (_time.sub("T", page_new.replace(html_new["ЗАТРУДНЕНИЕ"], "§"))
            == _time.sub("T", page_old.replace(html_old["ЗАТРУДНЕНИЕ"], "§")))
    assert all_new.replace(pptx_new["ПРОБЛЕМА"], "§") == all_old.replace(pptx_old["ПРОБЛЕМА"], "§")


def test_h4_helper_contract():
    from engines.channel_action import drop_named_hero
    assert drop_named_hero(["A", "B"], "A", True) == ["B"]
    assert drop_named_hero(["A", "B"], "A", False) == ["A", "B"]
    assert drop_named_hero(["A", "B"], None, True) == ["A", "B"]
    src = ["A", "B"]
    assert drop_named_hero(src, "Z", True) == src and drop_named_hero(src, "Z", True) is not src


# ═══ H5: оси столбчатых диаграмм ════════════════════════════════════════════

_BAR_COLUMN = {
    getattr(XL_CHART_TYPE, n) for n in dir(XL_CHART_TYPE)
    if re.match(r"^(BAR|COLUMN)_", n)
}


@pytest.mark.parametrize("name", ["fixture", "full"])
def test_h5_bar_and_column_axes_start_at_zero_and_cover_data(name, decks):
    prs = Presentation(decks[name])
    checked = 0
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not (getattr(sh, "has_chart", False) and sh.has_chart):
                continue
            ch = sh.chart
            if ch.chart_type not in _BAR_COLUMN:
                continue
            vals = [v for pl in ch.plots for se in pl.series for v in se.values if v is not None]
            if not vals or min(vals) < 0:
                continue
            axis = ch.value_axis
            assert axis.minimum_scale in (None, 0), (si, ch.chart_type, axis.minimum_scale)
            assert axis.maximum_scale is None or axis.maximum_scale >= max(vals), (
                si, axis.maximum_scale, max(vals))
            checked += 1
    assert checked >= 1, "в колоде нет ни одной столбчатой диаграммы с неотрицательными данными"
